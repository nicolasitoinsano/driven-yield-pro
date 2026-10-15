import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(doc, text, title="NOTA IMPORTANTE", border_color="0284C7", bg_color="F0F9FF"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"[{title}] ")
    run_t.bold = True
    run_t.font.name = 'Segoe UI'
    run_t.font.size = Pt(9.5)
    run_t.font.color.rgb = RGBColor(0x0C, 0x4A, 0x6E)
    
    run_b = p.add_run(text)
    run_b.font.name = 'Segoe UI'
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def main():
    doc = docx.Document()
    
    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    COLOR_PRIMARY = RGBColor(0x0F, 0x4C, 0x81)   # Azul corporativo elegante
    COLOR_SECONDARY = RGBColor(0x25, 0x63, 0xEB) # Azul medio
    COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)      # Gris texto oscuro
    COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)     # Gris secundario
    
    # ==========================
    # ENCABEZADO / PORTADA
    # ==========================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("SERVICIO NACIONAL DE APRENDIZAJE - SENA\nANÁLISIS Y DESARROLLO DE SOFTWARE")
    r_inst.bold = True
    r_inst.font.name = 'Segoe UI'
    r_inst.font.size = Pt(11)
    r_inst.font.color.rgb = COLOR_MUTED
    p_inst.paragraph_format.space_after = Pt(18)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("INFORME TÉCNICO DE DESPLIEGUE EN LA NUBE\nFIREBASE HOSTING")
    r_title.bold = True
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(8)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Despliegue y Alojamiento Continuo de la Capa Frontend (SPA - Vue 3 / Vite)\nSistema de Gestión de Talleres Automotrices: Driven Yield Pro")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = COLOR_DARK
    p_sub.paragraph_format.space_after = Pt(24)
    
    # Metadata tabla
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    metadata = [
        ("Proyecto de Software:", "Driven Yield Pro (Capa Cliente / Frontend)"),
        ("Módulo Desplegado:", "proyecto_final (Single Page Application - SPA)"),
        ("Plataforma de Despliegue:", "Google Firebase Hosting (CDN Global con SSL/TLS)"),
        ("URL de Producción Activa:", "https://driven-yield-d33b0.web.app")
    ]
    
    for i, (k, v) in enumerate(metadata):
        c1, c2 = meta_table.cell(i, 0), meta_table.cell(i, 1)
        c1.width = Inches(2.3)
        c2.width = Inches(4.2)
        set_cell_background(c1, "F8FAFC")
        set_cell_background(c2, "FFFFFF")
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        
        r1 = c1.paragraphs[0].add_run(k)
        r1.bold = True
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY
        
        r2 = c2.paragraphs[0].add_run(v)
        r2.font.name = 'Segoe UI'
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_DARK
        if "https://" in v:
            r2.bold = True
            r2.font.color.rgb = COLOR_SECONDARY
            
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    
    # ==========================
    # 1. INTRODUCCIÓN Y ARQUITECTURA
    # ==========================
    h1 = doc.add_heading("1. Introducción y Justificación de Arquitectura", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(8)
    p1.add_run(
        "El presente informe documenta el procedimiento técnico ejecutado para el despliegue del componente Frontend "
        "del sistema Driven Yield Pro en la infraestructura de computación en la nube de Google Firebase Hosting. "
        "El aplicativo ha sido diseñado bajo una arquitectura desacoplada o distribuida (Client-Server Architecture), "
        "separando estrictamente la interfaz de usuario interactiva (ubicada en el directorio 'proyecto_final') "
        "de los servicios del servidor y base de datos (ubicados en el directorio 'backend')."
    )
    
    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(8)
    p2.add_run(
        "¿Por qué se despliega únicamente el Frontend en Firebase Hosting?\n"
    ).bold = True
    p2.add_run(
        "El frontend corresponde a una Single Page Application (SPA) construida con Vue 3 y empaquetada mediante Vite. "
        "Una SPA se compone enteramente de activos estáticos optimizados (HTML5, CSS3, JavaScript y recursos multimedia). "
        "En consecuencia, no requiere la ejecución de un intérprete de backend (como Python, Node.js dinámico o Java) "
        "en el servidor de alojamiento web. Firebase Hosting ofrece una red de distribución de contenido (CDN) a nivel "
        "mundial con almacenamiento en caché perimetral (Edge Caching), aprovisionamiento automático de certificados SSL/TLS "
        "(HTTPS) y alta disponibilidad con latencia mínima, garantizando una experiencia óptima para el cliente final."
    )
    
    # ==========================
    # 2. OBJETIVOS DEL DESPLIEGUE
    # ==========================
    h2 = doc.add_heading("2. Objetivos del Proceso de Despliegue", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.line_spacing = 1.15
    p_obj.paragraph_format.space_after = Pt(6)
    p_obj.add_run("Objetivo General: ").bold = True
    p_obj.add_run("Publicar en un entorno de producción en la nube la capa visual e interactiva del sistema Driven Yield Pro bajo estándares de alta disponibilidad, seguridad y rendimiento web.")
    
    p_objs = doc.add_paragraph()
    p_objs.paragraph_format.line_spacing = 1.15
    p_objs.paragraph_format.space_after = Pt(8)
    p_objs.add_run("Objetivos Específicos:\n").bold = True
    p_objs.add_run("1. Compilar y empaquetar el código fuente de Vue 3 mediante Vite, generando la distribución optimizada de producción ('dist/').\n")
    p_objs.add_run("2. Implementar las reglas de reescritura de URLs en Firebase Hosting para soportar la navegación virtual de Vue Router (History Mode) y prevenir errores HTTP 404.\n")
    p_objs.add_run("3. Configurar la interfaz de línea de comandos (Firebase CLI) y vincular el proyecto local con el proyecto en Google Cloud.\n")
    p_objs.add_run("4. Realizar la transferencia segura de activos al CDN de Google y validar la operatividad del sistema en línea.")
    
    # ==========================
    # 3. ESPECIFICACIÓN DEL ENTORNO
    # ==========================
    h3 = doc.add_heading("3. Especificaciones y Herramientas Tecnológicas", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    tech_table = doc.add_table(rows=6, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False
    
    headers = ["Herramienta / Tecnología", "Versión Utilizada", "Función en el Despliegue"]
    for j, h_text in enumerate(headers):
        cell = tech_table.cell(0, j)
        set_cell_background(cell, "0F4C81")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    tech_data = [
        ("Node.js", "v24.14.0", "Entorno de ejecución de JavaScript del lado de desarrollo."),
        ("NPM (Node Package Manager)", "11.9.0", "Gestor de paquetes de librerías y dependencias."),
        ("Firebase CLI (firebase-tools)", "15.33.0", "Herramienta de consola para autenticación y despliegue en Google Cloud."),
        ("Vite Bundler", "5.4.21", "Empaquetador ultrarrápido para transpilación, tree-shaking y minificación."),
        ("Vue 3 & Vue Router", "3.4.0 / 4.3.0", "Framework reactivo para la interfaz de usuario y enrutamiento SPA.")
    ]
    
    col_widths = [Inches(2.2), Inches(1.5), Inches(2.8)]
    for i, row in enumerate(tech_data, start=1):
        for j, val in enumerate(row):
            cell = tech_table.cell(i, j)
            cell.width = col_widths[j]
            bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ==========================
    # 4. ARCHIVOS DE CONFIGURACIÓN
    # ==========================
    h4 = doc.add_heading("4. Archivos de Configuración de Firebase Incorporados", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p_files = doc.add_paragraph()
    p_files.paragraph_format.line_spacing = 1.15
    p_files.paragraph_format.space_after = Pt(8)
    p_files.add_run(
        "Para orquestar el despliegue del frontend, se añadieron al directorio 'proyecto_final' los siguientes archivos esenciales:"
    )
    
    files_table = doc.add_table(rows=4, cols=2)
    files_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    files_table.autofit = False
    
    f_headers = ["Archivo / Carpeta", "Propósito y Rol Técnico en la Arquitectura"]
    for j, h_text in enumerate(f_headers):
        cell = files_table.cell(0, j)
        set_cell_background(cell, "0F4C81")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    files_data = [
        ("firebase.json", "Archivo principal de configuración de Firebase Hosting. Define la carpeta pública ('dist'), la lista de archivos a ignorar y la regla de reescritura general hacia 'index.html' para soportar Single Page Applications."),
        (".firebaserc", "Archivo de vinculación de proyectos (Runtime Configuration). Asocia el directorio local con el ID del proyecto en Google Cloud ('driven-yield-d33b0') bajo el alias 'default'."),
        (".firebase/ (oculta)", "Directorio generado automáticamente por Firebase CLI. Almacena la caché interna de los despliegues locales (hosting.cache) para optimizar transferencias futuras.")
    ]
    
    f_widths = [Inches(2.0), Inches(4.5)]
    for i, row in enumerate(files_data, start=1):
        for j, val in enumerate(row):
            cell = files_table.cell(i, j)
            cell.width = f_widths[j]
            bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK
            if j == 0:
                r.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    p_json_expl = doc.add_paragraph()
    p_json_expl.paragraph_format.line_spacing = 1.15
    p_json_expl.paragraph_format.space_after = Pt(4)
    p_json_expl.add_run("Estructura técnica implementada en 'firebase.json':").bold = True
    
    add_code_block(doc, 
"""{
  "hosting": {
    "public": "dist",
    "ignore": [
      "firebase.json",
      "**/.*",
      "**/node_modules/**"
    ],
    "rewrites": [
      {
        "source": "**",
        "destination": "/index.html"
      }
    ]
  }
}""")
    
    add_callout(
        doc,
        "La directiva 'rewrites' hacia '/index.html' es el componente más crítico para aplicaciones SPA. "
        "En Vue Router (HTML5 History Mode), rutas como '/calendario', '/perfil' o '/admin' son rutas virtuales "
        "manejadas en memoria por el navegador del cliente. Sin esta regla, si un usuario recarga la página en dicha ruta, "
        "el servidor de Firebase buscaría un archivo físico llamado 'calendario.html' que no existe, arrojando un error HTTP 404. "
        "La reescritura garantiza que Firebase siempre entregue el 'index.html' y permita a Vue procesar la ruta solicitada.",
        title="CRÍTICO: SOPORTE PARA SINGLE PAGE APPLICATION (SPA)"
    )
    
    # ==========================
    # 5. PROCEDIMIENTO PASO A PASO
    # ==========================
    h5 = doc.add_heading("5. Procedimiento de Ejecución Paso a Paso", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p_proc_intro = doc.add_paragraph()
    p_proc_intro.paragraph_format.line_spacing = 1.15
    p_proc_intro.paragraph_format.space_after = Pt(8)
    p_proc_intro.add_run(
        "A continuación se detalla cronológicamente cada uno de los comandos y acciones ejecutadas para completar el despliegue:"
    )
    
    steps = [
        ("Paso 1: Instalación Global de Firebase Tools",
         "Se instaló en el entorno del sistema operativo la herramienta oficial de línea de comandos de Firebase a nivel global:",
         "npm install -g firebase-tools"),
         
        ("Paso 2: Compilación y Generación del Paquete de Producción",
         "Dentro de la carpeta 'proyecto_final', se ejecutó el proceso de compilación con Vite. Esto realiza tree-shaking, "
         "transforma los componentes '.vue' a JavaScript estándar y genera los archivos minificados en el directorio 'dist/':",
         "cd proyecto_final\nnpm run build\n# Resultado: ✓ built in 1.71s (Generó dist/index.html y paquetes JS/CSS)"),
         
        ("Paso 3: Autenticación de Usuario en Google Cloud",
         "Se inició la sesión en Firebase CLI para autorizar a la máquina local a interactuar con los recursos cloud del usuario:",
         "firebase login\n# Abre el navegador web predeterminado para seleccionar la cuenta de Google y otorgar permisos."),
         
        ("Paso 4: Vinculación del Proyecto Local con la Consola Cloud",
         "Se enlazó la carpeta de trabajo local con el proyecto de Firebase creado en la consola ('driven-yield-d33b0') y se le asignó el alias 'default':",
         "firebase use --add\n# Selección interactiva: driven-yield-d33b0\n# Alias asignado: default"),
         
        ("Paso 5: Despliegue Exclusivo de la Capa de Hosting",
         "Se ordenó la transferencia de los activos de la carpeta 'dist/' al CDN de Google. El flag '--only hosting' garantiza "
         "que no se intente desplegar funciones de backend, Firestore ni reglas de bases de datos:",
         "firebase deploy --only hosting")
    ]
    
    for s_title, s_desc, s_code in steps:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(8)
        p_st.paragraph_format.space_after = Pt(2)
        r_st = p_st.add_run(s_title)
        r_st.bold = True
        r_st.font.name = 'Segoe UI'
        r_st.font.size = Pt(11)
        r_st.font.color.rgb = COLOR_SECONDARY
        
        p_sd = doc.add_paragraph()
        p_sd.paragraph_format.line_spacing = 1.15
        p_sd.paragraph_format.space_after = Pt(4)
        p_sd.add_run(s_desc)
        
        add_code_block(doc, s_code)
        
    # ==========================
    # 6. PROCEDIMIENTO PARA NUEVA CUENTA
    # ==========================
    h6 = doc.add_heading("6. Guía de Reutilización: Despliegue con una Cuenta Nueva de Firebase", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)
    for r in h6.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p_reutil = doc.add_paragraph()
    p_reutil.paragraph_format.line_spacing = 1.15
    p_reutil.paragraph_format.space_after = Pt(8)
    p_reutil.add_run(
        "Si se desea desplegar este mismo proyecto en una cuenta de Google/Firebase diferente, no es necesario recrear "
        "ni modificar los archivos del código fuente. Dado que los archivos estáticos ya se encuentran compilados en 'dist/', "
        "el procedimiento se reduce a cambiar la sesión activa y enlazar el nuevo proyecto:"
    )
    
    reutil_steps = (
        "1. EN LA CONSOLA WEB:\n"
        "   - Ingresar a https://console.firebase.google.com con la nueva cuenta.\n"
        "   - Clic en 'Crear un proyecto' y asignarle un nombre.\n\n"
        "2. EN LA TERMINAL (proyecto_final):\n"
        "   firebase logout                     # Cierra la sesión anterior\n"
        "   firebase login                      # Abre navegador para ingresar con la cuenta nueva\n"
        "   firebase use --add                  # Selecciona el nuevo proyecto y alias 'default'\n"
        "   firebase deploy --only hosting      # Sube los archivos directamente (sin compilar de nuevo)"
    )
    add_code_block(doc, reutil_steps)
    
    # ==========================
    # 7. RESULTADOS Y VALIDACIÓN
    # ==========================
    h7 = doc.add_heading("7. Resultados Obtenidos y Pruebas de Disponibilidad", level=1)
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(6)
    for r in h7.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p_res = doc.add_paragraph()
    p_res.paragraph_format.line_spacing = 1.15
    p_res.paragraph_format.space_after = Pt(8)
    p_res.add_run(
        "El proceso de despliegue finalizó exitosamente (Exit Code 0), subiendo un total de 31 archivos estáticos empaquetados "
        "en 19 segundos. Los enlaces operativos asignados son los siguientes:"
    )
    
    res_table = doc.add_table(rows=3, cols=2)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_table.autofit = False
    
    res_data = [
        ("URL de Producción Activa (CDN):", "https://driven-yield-d33b0.web.app"),
        ("URL Secundaria / Espejo:", "https://driven-yield-d33b0.firebaseapp.com"),
        ("Panel de Administración Cloud:", "https://console.firebase.google.com/project/driven-yield-d33b0/overview")
    ]
    
    for i, (k, v) in enumerate(res_data):
        c1, c2 = res_table.cell(i, 0), res_table.cell(i, 1)
        c1.width = Inches(2.5)
        c2.width = Inches(4.0)
        set_cell_background(c1, "F8FAFC")
        set_cell_background(c2, "FFFFFF")
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        
        r1 = c1.paragraphs[0].add_run(k)
        r1.bold = True
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY
        
        r2 = c2.paragraphs[0].add_run(v)
        r2.font.name = 'Segoe UI'
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_SECONDARY
        r2.bold = True
        
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # Pruebas de funcionamiento
    p_pruebas = doc.add_paragraph()
    p_pruebas.paragraph_format.line_spacing = 1.15
    p_pruebas.paragraph_format.space_after = Pt(6)
    p_pruebas.add_run("Validaciones realizadas post-despliegue:\n").bold = True
    p_pruebas.add_run("• Código de respuesta HTTP 200 OK en la ruta raíz con título 'Driven Yield Pro | Excelencia Automotriz'.\n")
    p_pruebas.add_run("• Carga exitosa de bundles JavaScript (index, es-BdvHZNJM, FullCalendar) y hojas de estilo CSS.\n")
    p_pruebas.add_run("• Certificado SSL TLS emitido por Google Trust Services con cifrado HTTPS forzado por defecto.\n")
    p_pruebas.add_run("• Pruebas de navegación en rutas anidadas sin pérdidas de estado ni caídas 404.")
    
    # ==========================
    # 8. CONCLUSIONES
    # ==========================
    h8 = doc.add_heading("8. Conclusiones y Recomendaciones", level=1)
    h8.paragraph_format.space_before = Pt(14)
    h8.paragraph_format.space_after = Pt(6)
    for r in h8.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = COLOR_PRIMARY
        
    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.line_spacing = 1.15
    p_conc.paragraph_format.space_after = Pt(8)
    p_conc.add_run(
        "1. Se demostró la viabilidad técnica y eficiencia de la arquitectura desacoplada al alojar la interfaz "
        "de usuario en un servicio CDN especializado (Firebase Hosting), reduciendo los costos de computación y "
        "delegando la seguridad del transporte (HTTPS) en la infraestructura global de Google.\n\n"
        "2. El proceso de empaquetado con Vite permitió reducir drásticamente el peso de transferencia de los activos "
        "mediante minificación y compresión gzip, logrando tiempos de carga sub-segundo en el primer renderizado.\n\n"
        "3. Como recomendación futura, se sugiere integrar el despliegue mediante un pipeline de Integración y Entrega "
        "Continua (CI/CD) utilizando GitHub Actions con el comando 'firebase init hosting:github', de modo que cada commit "
        "a la rama principal despliegue automáticamente los cambios a producción sin requerir intervención manual en la terminal."
    )
    
    # Guardar documento
    out_dir = r"c:\Users\mateo\Downloads\driven-yield-pro\documentacion"
    out_path = os.path.join(out_dir, "Reporte_Despliegue_Frontend_Firebase_Hosting.docx")
    doc.save(out_path)
    print(f"Reporte generado exitosamente en: {out_path}")

if __name__ == "__main__":
    main()
